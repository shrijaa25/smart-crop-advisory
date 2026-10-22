"""Load SmartCrop rules from the two Excel workbooks.

This module reads source data only. It does not activate unverified rules.
Requires: openpyxl
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]

WEATHER_WORKBOOK = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "SmartCrop_RAW_Prevention_Temperature_Symptoms_Updated.xlsx"
)

FERTILIZER_WORKBOOK = (
    PROJECT_ROOT
    / "data"
    / "fertilizer"
    / "SmartCrop_Tomato_Chilli_Fertilizer_Schedule.xlsx"
)


def _read_sheet(
    path: Path,
    sheet_name: str,
) -> list[dict[str, Any]]:
    if not path.exists():
        raise FileNotFoundError(
            f"Knowledge workbook not found: {path}"
        )

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_name}' not found in {path.name}. "
                f"Available sheets: {workbook.sheetnames}"
            )

        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return []

        column_names = [
            str(header).strip()
            if header is not None
            else f"column_{index + 1}"
            for index, header in enumerate(headers)
        ]

        result = []

        for row in rows:
            if not any(value is not None for value in row):
                continue

            result.append({
                column_names[index]: (
                    row[index] if index < len(row) else None
                )
                for index in range(len(column_names))
            })

        return result

    finally:
        workbook.close()


def load_weather_disease_rules(
    path: Path | str = WEATHER_WORKBOOK,
) -> list[dict[str, Any]]:
    """Load weather-risk reference rows."""
    return _read_sheet(
        Path(path),
        "WEATHER_DISEASE_RISK",
    )


def load_day_1_30_tasks(
    path: Path | str = WEATHER_WORKBOOK,
) -> list[dict[str, Any]]:
    """Load the first 30 days of task reference rows."""
    return _read_sheet(
        Path(path),
        "DAY_1_30_TASKS",
    )


def load_drench_day_1_30_rules(
    path: Path | str = WEATHER_WORKBOOK,
) -> list[dict[str, Any]]:
    """Load source drench rows for review-only display."""
    return _read_sheet(
        Path(path),
        "DRENCH_DAY_1_30",
    )


def load_preplant_tasks(
    path: Path | str = WEATHER_WORKBOOK,
) -> list[dict[str, Any]]:
    """Load pre-plant task reference rows."""
    return _read_sheet(
        Path(path),
        "PREPLANT_TASKS",
    )


def load_fertilizer_rules(
    path: Path | str = FERTILIZER_WORKBOOK,
) -> dict[str, list[dict[str, Any]]]:
    """Load tomato and chilli fertilizer reference rows."""
    return {
        "chilli": _read_sheet(
            Path(path),
            "CHILLI_FERTILIZER",
        ),
        "tomato": _read_sheet(
            Path(path),
            "TOMATO_FERTILIZER",
        ),
        "task_windows": _read_sheet(
            Path(path),
            "TASK_WINDOWS",
        ),
    }


def is_rule_eligible(row: dict[str, Any]) -> bool:
    """Only explicitly active, verified rules are eligible."""
    status = str(
        row.get("review_status") or ""
    ).strip().lower()

    active_value = row.get("is_active")

    active = (
        active_value is True
        or str(active_value).strip().lower()
        in {"true", "1", "yes"}
    )

    return active and status in {"verified", "approved"}


def load_raw_prevention(path: Path | str = WEATHER_WORKBOOK):
    """Load the original prevention reference records."""
    return _read_sheet(Path(path), "RAW_Prevention")


def load_raw_pesticide_disease_rules(
    path: Path | str = WEATHER_WORKBOOK,
):
    """Load disease-control source records for review."""
    return _read_sheet(Path(path), "RAW_PesticideDiseases")


def load_raw_pesticide_pest_rules(
    path: Path | str = WEATHER_WORKBOOK,
):
    """Load pest-control source records for review."""
    return _read_sheet(Path(path), "RAW_PesticidePests")


def load_source_rules_review(path: Path | str = WEATHER_WORKBOOK):
    """Load source rules together with their review and activation status."""
    return _read_sheet(Path(path), "SOURCE_RULES_REVIEW")


def load_cnn_data_requirements(path: Path | str = WEATHER_WORKBOOK):
    """Load the image-model dataset requirements."""
    return _read_sheet(Path(path), "CNN_DATA_REQUIREMENTS")
