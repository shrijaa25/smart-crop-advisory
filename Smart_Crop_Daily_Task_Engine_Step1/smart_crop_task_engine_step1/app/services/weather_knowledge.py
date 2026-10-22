from __future__ import annotations

import re
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[2]

WORKBOOK_PATH = (
    PROJECT_ROOT
    / "data"
    / "raw"
    / "SmartCrop_RAW_Prevention_Temperature_Symptoms_Updated.xlsx"
)


def check_workbook() -> Path:
    """Return the workbook path after verifying it exists."""
    if not WORKBOOK_PATH.is_file():
        raise FileNotFoundError(
            f"Weather knowledge workbook not found: {WORKBOOK_PATH}"
        )
    return WORKBOOK_PATH


def get_workbook_sheets() -> list[str]:
    """Return the worksheet names from the configured workbook."""
    workbook = load_workbook(
        check_workbook(),
        read_only=True,
        data_only=True,
    )
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def _text(value: Any) -> str:
    """Convert an Excel cell value to clean text."""
    if value is None:
        return ""
    return str(value).strip()


def _normalise(value: Any) -> str:
    return _text(value).casefold().strip()


def _read_sheet(sheet_name: str) -> list[dict[str, Any]]:
    """Read a sheet, cached until the workbook file changes on disk.

    Callers must treat the returned rows as read-only.
    """
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError(
            f"Knowledge workbook not found: {WORKBOOK_PATH}"
        )
    return _read_sheet_cached(sheet_name, WORKBOOK_PATH.stat().st_mtime_ns)


@lru_cache(maxsize=32)
def _read_sheet_cached(sheet_name: str, _mtime_ns: int) -> list[dict[str, Any]]:
    """Read Excel rows while preserving duplicate column headings."""

    workbook = load_workbook(
        WORKBOOK_PATH,
        read_only=True,
        data_only=True,
    )

    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_name}' not found. "
                f"Available sheets: {workbook.sheetnames}"
            )

        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)

        original_headers = [
            _text(value) for value in next(rows, ())
        ]

        # Keep repeated headings instead of silently overwriting them.
        counts = {}
        headers = []

        for header in original_headers:
            if not header:
                headers.append("")
                continue

            counts[header] = counts.get(header, 0) + 1
            suffix = f".{counts[header] - 1}" if counts[header] > 1 else ""
            headers.append(f"{header}{suffix}")

        records = []

        for row in rows:
            record = {
                headers[i]: value
                for i, value in enumerate(row)
                if i < len(headers) and headers[i]
            }

            if any(value is not None for value in record.values()):
                records.append(record)

        return records
    finally:
        workbook.close()


def _weather_value(weather: dict, *names: str) -> float | None:
    """Read a numeric weather value from common API field names."""
    for name in names:
        value = weather.get(name)

        if value is None:
            continue

        try:
            return float(value)
        except (TypeError, ValueError):
            continue

    return None


def _crop_matches(crop_scope: Any, crop: str) -> bool:
    """Match a crop without treating unrelated crops as equivalent."""
    scope = _normalise(crop_scope)
    crop_name = _normalise(crop)

    if not scope or not crop_name:
        return False

    if crop_name in scope:
        return True

    # These are explicit multi-crop phrases in the workbook.
    return (
        "tomato and chilli" in scope
        or "tomato/chilli" in scope
        or "chilli and tomato" in scope
    ) and crop_name in {"tomato", "chilli"}


def _profile_matches(
    profile_text: str,
    weather: dict,
) -> bool | None:
    """Match only explicit AIR-temperature conditions from workbook text.

    This project feature is temperature-based. Humidity, rainfall, and soil
    moisture text is intentionally ignored. For a stated temperature range,
    compare it with the day's observed/forecast min-max interval.

    Returns True for a match, False for a definite mismatch, and None when
    no usable numeric air-temperature rule or weather data is available.
    Soil-temperature rules are never compared with air temperature.
    """
    text = _normalise(profile_text)
    observed = observed_temperature_range(weather)
    if observed is None:
        return None
    low, high = observed

    checks: list[bool] = []

    def is_soil_temperature(match_start: int) -> bool:
        context = text[max(0, match_start - 48):match_start]
        tail = re.split(r"[.;,]", context)[-1]
        return bool(
            re.search(
                r"\bsoil(?:\s+(?:temperature|temp))?\s*"
                r"(?:(?:around|about|at)\s*)?$",
                tail,
            )
        )

    # Explicit comparisons: temperature >30°C, temp <= 18 C, etc.
    comparison_pattern = (
        r"(?:temperature|temp)\s*(>=|<=|>|<)\s*"
        r"(-?\d+(?:\.\d+)?)\s*°?\s*c"
    )
    for match in re.finditer(comparison_pattern, text):
        if is_soil_temperature(match.start()):
            continue
        operator = match.group(1)
        threshold = float(match.group(2))
        if operator in {">", ">="}:
            checks.append(high > threshold if operator == ">" else high >= threshold)
        else:
            checks.append(low < threshold if operator == "<" else low <= threshold)

    # Explicit ranges: 25–30°C, temperature 25-30 C, etc.
    range_pattern = (
        r"(?:(?:air\s+)?(?:temperature|temp)\s*)?"
        r"(-?\d+(?:\.\d+)?)\s*°?\s*"
        r"[-–—]\s*(-?\d+(?:\.\d+)?)\s*°?\s*c"
    )
    for match in re.finditer(range_pattern, text):
        if is_soil_temperature(match.start()):
            continue
        rule_low, rule_high = sorted(
            (float(match.group(1)), float(match.group(2)))
        )
        # The day's interval and the rule interval must overlap.
        checks.append(low <= rule_high and high >= rule_low)

    if not checks:
        return None

    return all(checks)


def _compare(value: float, operator: str, threshold: float) -> bool:
    if operator == ">":
        return value > threshold
    if operator == "<":
        return value < threshold
    if operator == ">=":
        return value >= threshold
    if operator == "<=":
        return value <= threshold
    return False


def _qualitative_weather_candidate(text: str, weather: dict) -> bool:
    """Match RAW_Prevention only when explicit numeric weather rules match."""
    return _profile_matches(_text(text), weather) is True


def normalise_weather_for_advisory(weather: dict, target_date: str) -> dict:
    """Use same-day current temperature or the selected date daily maximum."""
    if not isinstance(weather, dict):
        weather = {}

    forecast = weather.get("forecast") or {}
    current = weather.get("current") or {}
    observed_at = str(weather.get("observed_at") or "")
    same_day_observation = (
        observed_at[:10] == target_date
        and _normalise(weather.get("observation_type")) == "observed"
        and _normalise(weather.get("status")) == "available"
    )

    current_temp = _weather_value(current, "temperature_2m", "temperature_c")
    forecast_max = _weather_value(
        forecast, "temperature_max_c", "temperature_2m_max"
    )
    weather_type = _normalise(weather.get("weather_type"))
    if target_date == observed_at[:10] and same_day_observation:
        temperature = current_temp
        temperature_source = "current_temperature"
    else:
        temperature = forecast_max
        temperature_source = (
            "historical_daily_max"
            if weather_type == "historical"
            else "forecast_daily_max"
        )

    humidity = (
        _weather_value(
            current, "relative_humidity_2m", "relative_humidity", "humidity"
        )
        if same_day_observation else None
    )
    rainfall = _weather_value(
        forecast, "precipitation_sum_mm", "precipitation_sum", "rainfall_mm"
    )
    return {
        "status": weather.get("status"),
        "target_date": target_date,
        "weather_type": weather.get("weather_type"),
        "temperature_c": temperature,
        "relative_humidity": humidity,
        "precipitation_mm": rainfall,
        "temperature_source": temperature_source if temperature is not None else "unavailable",
        "humidity_source": (
            "same_day_current_observation" if humidity is not None
            else "unavailable"
        ),
        "rainfall_source": (
            "forecast_daily_sum" if rainfall is not None else "unavailable"
        ),
    }




def _pesticide_disease_details(crop: str, disease: str) -> dict[str, str]:
    """Look up a treatment row after a symptom assessment requests it."""
    wanted_crop = _normalise(crop)
    wanted_disease = _normalise(disease)

    for row in _read_sheet("RAW_PesticideDiseases"):
        row_crop = _normalise(row.get("Crop"))
        row_disease = _normalise(row.get("Disease / Pest"))
        crop_matches = _crop_matches(row_crop, wanted_crop)
        disease_matches = (
            wanted_disease == row_disease
            or wanted_disease in row_disease
        )
        if not (crop_matches and disease_matches):
            continue

        return {
            "disease": _text(row.get("Disease / Pest")),
            "marathi_name": _text(row.get("Disease Name (Marathi)")),
            "chemical_name": _text(row.get("Pesticide (Chemical Name)")),
            "combination": _text(row.get("Combination")),
            "brand_name": _text(row.get("Brand Name")),
            "dosage_instructions": _text(
                row.get("Dosage & Application Instructions")
            ),
            "water_ph": _text(row.get("Water pH")),
            "weather_conditions": _text(row.get("Weather Conditions")),
            "pre_harvest_interval": _text(row.get("Pre-Harvest Interval")),
            "source_sheet": "RAW_PesticideDiseases",
            "review_status": "needs_review",
        }

    return {
        "source_sheet": "RAW_PesticideDiseases",
        "review_status": "unavailable",
        "message": "No exact crop-and-disease treatment match was found in the workbook.",
    }


def _raw_prevention_temperature_match(
    crop: str,
    temperature: float,
) -> list[dict[str, Any]]:
    """Return RAW_Prevention rows whose numeric temperature rule matches."""
    crop_name = _normalise(crop)
    matches = []
    seen = set()

    for row in _read_sheet("RAW_Prevention"):
        if _normalise(row.get("Crop")) != crop_name:
            continue

        disease = _text(row.get("Disease"))
        minimum = _weather_value(row, "Temperature Min (°C)")
        maximum = _weather_value(row, "Temperature Max (°C)")
        rule = _normalise(row.get("Temperature Rule"))

        if not disease or minimum is None and maximum is None:
            continue
        if any(
            term in _normalise(disease)
            for term in ("sticker", "spreader", "adjuvant")
        ):
            continue

        if minimum is not None and maximum is not None:
            matched = minimum <= temperature <= maximum
        elif minimum is not None and any(
            phrase in rule for phrase in ("and above", "or above", "at least")
        ):
            matched = temperature >= minimum
        elif maximum is not None and any(
            phrase in rule for phrase in ("below", "less than", "under")
        ):
            matched = temperature < maximum
        else:
            # A lone threshold without an explicit direction is ambiguous.
            continue

        if not matched:
            continue

        key = (crop_name, _normalise(disease))
        if key in seen:
            continue
        seen.add(key)

        matches.append({
            "crop": crop,
            "risk": disease,
            "marathi_name": _text(row.get("Disease Name (Marathi)")),
            "temperature_c": temperature,
            "temperature_min_c": minimum,
            "temperature_max_c": maximum,
            "temperature_rule": _text(row.get("Temperature Rule")),
            "temperature_operator": (
                "between_inclusive"
                if minimum is not None and maximum is not None
                else "greater_than_or_equal"
                if minimum is not None
                else "less_than"
            ),
            "pathogen_or_vector": _text(row.get("Pathogen Type")),
            "weather_trigger_as_supplied": _text(
                row.get("Primary Moisture / Weather Triggers")
            ),
            "symptoms_to_check": _text(
                row.get("Symptoms / Diagnostic Keys")
            ),
            "chemical_control": _text(row.get("Chemical Control")),
            "chemical_application_method": _text(
                row.get("Chemical Application Method")
            ),
            "chemical_dosage_instructions": _text(
                row.get("Chemical Control: Dosage & Instructions")
            ),
            "chemical_brand_name": _text(row.get("Chemical Brand Name")),
            "treatment_review_status": "needs_review",
            "diagnosis_confirmed": False,
            "chemical_advice_enabled": False,
            "match_status": "matched",
            "source_sheet": "RAW_Prevention",
        })

    return matches


_TEMPERATURE_OPERATORS = {"between_inclusive", "greater_than", "less_than"}


def observed_temperature_range(weather: dict) -> tuple[float, float] | None:
    """Return the (min, max) air temperature to compare with disease ranges.

    Uses the day's minimum and maximum when both are known. A single value
    (``temperature_c``) is treated as a zero-width range.
    """
    low = _weather_value(weather, "temperature_min_c")
    high = _weather_value(weather, "temperature_max_c")
    single = _weather_value(weather, "temperature_c", "temperature")

    if low is None and high is None:
        if single is None:
            return None
        return single, single
    if low is None:
        low = min(high, single) if single is not None else high
    if high is None:
        high = max(low, single) if single is not None else low
    return (low, high) if low <= high else (high, low)


def temperature_range_matches(row: dict, weather: dict) -> bool | None:
    """Compare the day's temperature range with a workbook row's range.

    The disease is considered possible when the observed [min, max] range
    overlaps the disease range from ``temperature_min_c`` /
    ``temperature_max_c`` / ``temperature_operator``:

    * between_inclusive: min <= observed_max and max >= observed_min
    * greater_than:      observed_max > min   (e.g. ">30 C")
    * less_than:         observed_min < max

    Returns None when the temperature or the row's thresholds are missing.
    """
    observed = observed_temperature_range(weather)
    if observed is None:
        return None
    low, high = observed

    operator = _normalise(row.get("temperature_operator"))
    row_min = _weather_value(row, "temperature_min_c")
    row_max = _weather_value(row, "temperature_max_c")

    if operator == "between_inclusive":
        if row_min is None or row_max is None:
            return None
        return row_min <= high and row_max >= low
    if operator == "greater_than":
        return None if row_min is None else high > row_min
    if operator == "less_than":
        return None if row_max is None else low < row_max
    return None


def describe_temperature_rule(row: dict) -> dict:
    """Describe a workbook row's temperature condition for display."""
    operator = _normalise(row.get("temperature_operator"))
    row_min = _weather_value(row, "temperature_min_c")
    row_max = _weather_value(row, "temperature_max_c")

    def fmt(value: float) -> str:
        return f"{value:g}"

    if operator == "between_inclusive" and None not in (row_min, row_max):
        label = f"{fmt(row_min)}\u2013{fmt(row_max)}\u00b0C"
    elif operator == "greater_than" and row_min is not None:
        label = f"above {fmt(row_min)}\u00b0C"
    elif operator == "less_than" and row_max is not None:
        label = f"below {fmt(row_max)}\u00b0C"
    else:
        label = ""
    return {
        "operator": operator or None,
        "min_c": row_min,
        "max_c": row_max,
        "label": label,
    }


def _structured_profile_matches(row: dict, weather: dict) -> bool | None:
    """Match only the temperature rule for a disease row.

    Humidity and moisture fields are deliberately not required because this
    project's alert feature is based on temperature only.
    """
    operator = _normalise(row.get("temperature_operator"))

    if operator in _TEMPERATURE_OPERATORS:
        return temperature_range_matches(row, weather)

    # Older workbook rows may store the numeric rule in descriptive text.
    return _profile_matches(
        _text(row.get("weather_profile_as_supplied")), weather
    )


def get_weather_advisory_candidates(
    crop: str,
    weather: dict,
) -> dict:
    """
    Match farm weather against the Excel knowledge base.

    This function creates risk/inspection candidates only.
    It does not diagnose a disease or prescribe a pesticide.
    """
    crop_name = _normalise(crop)

    if crop_name not in {"tomato", "chilli"}:
        return {
            "status": "unsupported_crop",
            "crop": crop,
            "risk_candidates": [],
            "prevention_candidates": [],
            "message": "This workbook matcher currently supports tomato and chilli.",
        }

    if not isinstance(weather, dict):
        weather = {}

    if _normalise(weather.get("status")) == "unavailable":
        return {
            "status": "weather_unavailable",
            "crop": crop,
            "risk_candidates": [],
            "prevention_candidates": [],
            "message": "Weather is unavailable. Use routine crop monitoring.",
        }

    temperature = _weather_value(weather, "temperature_c")
    if temperature is None:
        return {
            "status": "weather_unavailable",
            "crop": crop_name,
            "risk_candidates": [],
            "prevention_candidates": [],
            "message": "Current temperature is unavailable.",
        }

    prevention_candidates = []
    risk_candidates = _raw_prevention_temperature_match(
        crop_name,
        temperature,
    )

    return {
        "status": "ok",
        "crop": crop_name,
        "weather_used": {
            "temperature_c": temperature,
            # Retain these metadata fields for API compatibility; they are not
            # used by the temperature-only disease matching logic.
            "temperature_source": weather.get("temperature_source", "provided"),
        },
        "risk_candidates": risk_candidates,
        "prevention_candidates": prevention_candidates,
        "message": (
            "Weather matches are potential risks for inspection, "
            "not confirmed diagnoses or automatic chemical recommendations."
        ),
    }
